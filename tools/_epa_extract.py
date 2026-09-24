# tools/_epa_extract.py — Phase B extractor (called by tools/build-epa-master.js).
# Reads sources/EPA_Master/EPA_Master.xlsx and emits JSON on stdout:
# verification numbers + every row of the substances sheet.
import json
import openpyxl

wb = openpyxl.load_workbook('sources/EPA_Master/EPA_Master.xlsx', read_only=True)

# 1) verification numbers
verif = {}
for row in wb['أرقام التحقق'].iter_rows(min_row=2, values_only=True):
    if row[0]:
        verif[str(row[0]).strip()] = row[1]

# 2) the substances sheet
rows = []
it = wb['المواد الفعالة'].iter_rows(values_only=True)
next(it)  # header
for r in it:
    if r[0] is None and r[1] is None:
        continue
    rows.append({
        'pc_code': str(r[0] or '').strip(),
        'name': str(r[1] or '').replace('\xa0', ' ').strip(),
        'name_type': str(r[2] or ''),
        'name_note': str(r[3] or ''),
        'cas': str(r[4] or '').strip(),
        'cas_raw': str(r[5] or '').strip(),
        'cas_quality': str(r[6] or ''),
        'shared': str(r[7] or '').strip(),
        'shared_count': int(r[8] or 0),
        'active': int(r[9] or 0),
        'cancelled': int(r[10] or 0),
        'status_sheet': str(r[11] or ''),
        'cancel_reason': (str(r[12]).strip() if r[12] else ''),
        'cancel_reason_count': int(r[13] or 0)
    })

print(json.dumps({'verif': verif, 'rows': rows}, ensure_ascii=False))
